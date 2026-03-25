#!/usr/bin/env python3
"""
Fetch 50 diverse Kubernetes documentation pages from the official sitemap,
extract main article text, and write parallel .md, .txt, and .docx files.
"""

from __future__ import annotations

import html
import re
import ssl
import time
import urllib.request
import xml.etree.ElementTree as ET
from pathlib import Path

import certifi

from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt
import html2text

BASE = "https://kubernetes.io"
SITEMAP = f"{BASE}/en/sitemap.xml"
USER_AGENT = "Agentic-RAG-docs-mirror/1.0 (+local research)"
OUTPUT_ROOT = Path(__file__).resolve().parent.parent / "data"
TARGET_COUNT = 50
# Pull extra candidates so we still reach TARGET_COUNT if some pages fail.
SELECTION_POOL = 70
REQUEST_DELAY_SEC = 0.6

# robots.txt disallows most of /docs/reference/kubernetes-api/ — skip except allowlisted.
API_ALLOW = ("api-index", "labels-annotations-taints")


def fetch(url: str, timeout: int = 45) -> str:
    req = urllib.request.Request(
        url,
        headers={"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml"},
    )
    ctx = ssl.create_default_context(cafile=certifi.where())
    with urllib.request.urlopen(req, timeout=timeout, context=ctx) as resp:
        return resp.read().decode("utf-8", errors="replace")


def parse_sitemap_urls() -> list[str]:
    raw = fetch(SITEMAP)
    root = ET.fromstring(raw)
    ns = {"sm": "http://www.sitemaps.org/schemas/sitemap/0.9"}
    out: list[str] = []
    for loc in root.iter(f"{{{ns['sm']}}}loc"):
        if not loc.text:
            continue
        u = loc.text.strip().split('"')[0].rstrip("/")
        if "/docs/" not in u or "_print" in u:
            continue
        if "/docs/contribute/" in u:
            continue
        if "/docs/reference/kubernetes-api/" in u:
            if not any(a in u for a in API_ALLOW):
                continue
        if not u.startswith("https://kubernetes.io/docs"):
            continue
        out.append(u)
    # stable unique order
    seen: set[str] = set()
    uniq: list[str] = []
    for u in out:
        if u not in seen:
            seen.add(u)
            uniq.append(u)
    return uniq


def bucket(path: str) -> str:
    if path.startswith("/docs/concepts/"):
        return "concepts"
    if path.startswith("/docs/tasks/"):
        return "tasks"
    if path.startswith("/docs/tutorials/"):
        return "tutorials"
    if path.startswith("/docs/setup/"):
        return "setup"
    if path.startswith("/docs/reference/"):
        return "reference"
    if path.startswith("/docs/home"):
        return "home"
    return "other"


def select_diverse_urls(urls: list[str], n: int) -> list[str]:
    from urllib.parse import urlparse

    by_bucket: dict[str, list[str]] = {}
    for u in urls:
        p = urlparse(u).path or "/"
        b = bucket(p)
        by_bucket.setdefault(b, []).append(u)
    for b in by_bucket:
        by_bucket[b].sort()

    quotas = {
        "home": 1,
        "concepts": 18,
        "tasks": 16,
        "tutorials": 6,
        "setup": 5,
        "reference": 4,
    }
    chosen: list[str] = []
    home_url = f"{BASE}/docs/home/"
    if home_url.rstrip("/") in {x.rstrip("/") for x in urls}:
        chosen.append(home_url)
    elif f"{BASE}/docs/home" in urls:
        chosen.append(f"{BASE}/docs/home")

    def take(bucket_name: str, want: int) -> None:
        nonlocal chosen
        pool = [x for x in by_bucket.get(bucket_name, []) if x not in chosen]
        for x in pool[:want]:
            if len(chosen) >= n:
                return
            chosen.append(x)

    take("home", max(0, quotas["home"] - len([c for c in chosen if "home" in c])))
    for b in ("concepts", "tasks", "tutorials", "setup", "reference"):
        if len(chosen) >= n:
            break
        take(b, quotas[b])

    # fill remainder from any bucket
    if len(chosen) < n:
        rest = [u for u in urls if u not in chosen]
        for u in rest:
            if len(chosen) >= n:
                break
            chosen.append(u)

    out = chosen[:n]
    if not any("/docs/home" in c for c in out):
        rest = [x for x in out if "/docs/home" not in x]
        out = [f"{BASE}/docs/home/"] + rest[: max(0, n - 1)]
    return out[:n]


def slug_from_url(url: str) -> str:
    from urllib.parse import urlparse

    path = urlparse(url).path.strip("/")
    parts = [p for p in path.split("/") if p and p != "docs"]
    if not parts:
        return "home"
    s = "_".join(parts)
    s = re.sub(r"[^a-zA-Z0-9._-]+", "_", s).strip("_")
    return s[:120] or "page"


def extract_article(html_doc: str) -> tuple[str, str]:
    """Returns (title, inner_html_of_main_content)."""
    soup = BeautifulSoup(html_doc, "html.parser")
    title_el = soup.find("title")
    title = title_el.get_text(strip=True) if title_el else "Kubernetes Documentation"
    title = html.unescape(title.split("|")[0].strip() or title)

    main = soup.select_one("main#maindoc .td-content")
    if not main:
        main = soup.select_one("main#maindoc")
    if not main:
        main = soup.find("article")
    if not main:
        main = soup.find("body")
    assert main is not None
    # drop print-only / nav noise inside content if present
    for bad in main.select(".feedback, #print, .td-toc"):
        bad.decompose()
    inner = "".join(str(x) for x in main.children)
    return title, inner


def html_to_markdown(fragment_html: str) -> str:
    h = html2text.HTML2Text()
    h.ignore_links = False
    h.body_width = 0
    h.unicode_snob = True
    return h.handle(fragment_html).strip()


def html_to_plaintext(fragment_html: str) -> str:
    soup = BeautifulSoup(fragment_html, "html.parser")
    return soup.get_text("\n", strip=True)


def write_docx(path: Path, title: str, plaintext: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    doc.add_heading(title, level=0)
    for para in plaintext.split("\n\n"):
        p = para.strip()
        if p:
            doc.add_paragraph(p)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)

    all_urls = parse_sitemap_urls()
    picked = select_diverse_urls(all_urls, min(SELECTION_POOL, len(all_urls)))
    manifest: list[str] = []
    used_slugs: set[str] = set()
    tried: set[str] = set()

    def scrape_one(url: str, seq: int) -> bool:
        if url in tried:
            return False
        tried.add(url)
        slug = slug_from_url(url)
        if slug in used_slugs:
            slug = f"{slug}_{seq}"
        used_slugs.add(slug)
        print(f"[{len(manifest)+1}/{TARGET_COUNT}] {url}")
        try:
            page = fetch(url)
        except Exception as e:
            print(f"  skip (fetch error): {e}")
            return False
        try:
            title, inner = extract_article(page)
        except Exception as e:
            print(f"  skip (parse error): {e}")
            return False

        md_body = html_to_markdown(inner).strip()
        md_full = f"# {title}\n\n**Source:** {url}\n\n{md_body}\n"
        txt_body = html_to_plaintext(inner)
        txt_full = f"{title}\n{'=' * len(title)}\nSource: {url}\n\n{txt_body}\n"

        OUTPUT_ROOT.joinpath(f"{slug}.md").write_text(md_full, encoding="utf-8")
        OUTPUT_ROOT.joinpath(f"{slug}.txt").write_text(txt_full, encoding="utf-8")
        write_docx(OUTPUT_ROOT / f"{slug}.docx", title, txt_full)
        manifest.append(f"{slug}\t{url}\t{title}")
        time.sleep(REQUEST_DELAY_SEC)
        return True

    seq = 0
    for url in picked:
        if len(manifest) >= TARGET_COUNT:
            break
        seq += 1
        scrape_one(url, seq)

    if len(manifest) < TARGET_COUNT:
        for url in all_urls:
            if len(manifest) >= TARGET_COUNT:
                break
            if url in tried:
                continue
            seq += 1
            scrape_one(url, seq)

    (OUTPUT_ROOT / "manifest.tsv").write_text(
        "slug\turl\ttitle\n" + "\n".join(manifest) + "\n", encoding="utf-8"
    )
    print(f"Done. Wrote {len(manifest)} pages under {OUTPUT_ROOT}")


if __name__ == "__main__":
    main()
